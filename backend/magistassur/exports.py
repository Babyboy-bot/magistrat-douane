"""Fonctions d'export Excel et PDF pour les contrats Magistassur.

Les exports Excel et PDF sont générés à partir des MÊMES sections
(`_epargne_sections` / `_prevoyance_sections`) : impossible qu'ils divergent.
"""
import io
from datetime import date
from xml.sax.saxutils import escape as xml_escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from django.http import HttpResponse


# ── Helpers ──────────────────────────────────────────────────────────────────

def _fmt(value):
    if value is None or value == "":
        return "—"
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)


def _money(value):
    if value is None:
        return "—"
    try:
        return f"{int(value):,} FCFA".replace(",", " ")
    except Exception:
        return str(value)


_PERIODICITE_LABELS = {
    "mensuelle": "Mensuelle",
    "trimestrielle": "Trimestrielle",
    "semestrielle": "Semestrielle",
    "annuelle": "Annuelle",
}


def _beneficiaire_label(b):
    """Clause générique (conjoint / enfants) : pas de nom ni de taux."""
    if b.conjoint_corps:
        return "Conjoint(e) non séparé(e) de corps"
    if b.enfants_nes:
        return "Enfants nés ou à naître"
    return f"{b.nom} — {b.en_qualite_de} — {b.taux_repartition}%"


def _conseiller_rows(contrat):
    """Conseiller responsable du dossier (prise en charge prioritaire, sinon agent)."""
    agent = contrat.pris_en_charge_par or contrat.agent
    if not agent:
        return [("Conseiller", "Non assigné")]
    return [
        ("Conseiller", agent.full_name),
        ("Téléphone conseiller", _fmt(getattr(agent, "telephone", "") or "")),
    ]


# ── Sections communes Excel / PDF ────────────────────────────────────────────
# Chaque fonction retourne [(titre, [(label, valeur), ...]), ...]

def _epargne_sections(contrat):
    prime_annuelle = int(contrat.montant_cotisation or 0) * {
        "mensuelle": 12, "trimestrielle": 4, "semestrielle": 2,
    }.get(contrat.periodicite, 1)

    sections = [
        ("IDENTIFIANTS", [
            ("N° Police", _fmt(contrat.num_police) if contrat.num_police else "En attente"),
            ("Code Inspection", _fmt(contrat.code_inspection)),
            ("Statut", contrat.get_statut_display()),
            ("Date de souscription", _fmt(contrat.created_at.date() if contrat.created_at else None)),
        ]),
        ("SOUSCRIPTEUR", [
            ("Nom complet", f"{contrat.client.nom} {contrat.client.prenoms}"),
            ("Civilité", _fmt(contrat.civilite)),
            ("Date de naissance", _fmt(contrat.date_naissance)),
            ("Pays/Ville naissance", _fmt(contrat.pays_naissance)),
            ("NNI", _fmt(contrat.nni)),
            ("Type pièce", contrat.get_type_piece_display() if contrat.type_piece else "—"),
            ("N° pièce", _fmt(contrat.numero_piece)),
            ("Validité pièce", _fmt(contrat.validite_piece)),
            ("Situation matrimoniale", contrat.get_situation_matrimoniale_display() if contrat.situation_matrimoniale else "—"),
            ("Nom époux/épouse", _fmt(contrat.nom_epoux)),
            ("Adresse postale", _fmt(contrat.adresse_postale)),
            ("Téléphone", _fmt(contrat.telephone)),
            ("WhatsApp", _fmt(contrat.whatsapp)),
            ("Email personnel", _fmt(contrat.email_personnel)),
            ("Résidence principale", _fmt(contrat.residence_principale)),
            ("Statut professionnel", contrat.get_statut_professionnel_display() if contrat.statut_professionnel else "—"),
            ("Profession", _fmt(contrat.profession_detail)),
            ("Secteur d'activité", _fmt(contrat.secteur_activite)),
            ("Revenu mensuel estimé", _money(contrat.revenu_mensuel)),
            ("Origine revenu", contrat.get_origine_revenu_display() if contrat.origine_revenu else "—"),
        ]),
    ]

    # Assuré
    if contrat.assure_idem:
        assure_rows = [("Assuré", "Identique au souscripteur")]
    else:
        assure_rows = [
            ("Nom & Prénoms", f"{contrat.assure_nom} {contrat.assure_prenoms}"),
            ("Date de naissance", _fmt(contrat.assure_date_naissance)),
            ("Lieu de naissance", _fmt(contrat.assure_lieu_naissance)),
            ("NNI", _fmt(contrat.assure_nni)),
            ("Type pièce", _fmt(contrat.assure_type_piece)),
            ("N° pièce", _fmt(contrat.assure_numero_piece)),
            ("Lien avec souscripteur", _fmt(contrat.assure_lien)),
            ("Téléphone", _fmt(contrat.assure_telephone)),
            ("Email", _fmt(contrat.assure_email)),
            ("Profession", _fmt(contrat.assure_profession)),
            ("Déjà souscrit NSIA ?", _fmt(contrat.assure_nsia_anterieur)),
            ("N° police antérieur", _fmt(contrat.assure_nsia_num_police)),
        ]
    sections.append(("ASSURÉ", assure_rows))

    sections.append(("GARANTIES & COTISATIONS", [
        ("Durée (années)", f"{contrat.duree_annees} ans"),
        ("Périodicité", _PERIODICITE_LABELS.get(contrat.periodicite, contrat.periodicite)),
        ("Montant cotisation", _money(contrat.montant_cotisation)),
        ("Prime annuelle", _money(prime_annuelle)),
        ("Date 1ère cotisation", _fmt(contrat.date_premiere_cotisation)),
        ("Date d'effet", _fmt(contrat.date_effet)),
        ("Capital constitué", _money(contrat.capital_constitue)),
        ("Mode de paiement", _fmt(contrat.mode_paiement)),
        ("N° compte (RIB)", _fmt(contrat.numero_compte)),
        ("Frais de gestion", f"{contrat.frais_gestion_pct} %"),
        ("Frais accessoires", _money(contrat.frais_accessoires)),
        ("Coût police", _money(contrat.cout_police)),
        ("Taux d'intérêt", f"{contrat.taux_interet} %"),
    ]))

    sections.append(("PIÈCES JOINTES", [
        ("Pièce d'identité (recto)", "Fournie" if contrat.piece_identite else "NON FOURNIE"),
        ("Pièce d'identité (verso)", "Fournie" if contrat.piece_identite_verso else "NON FOURNIE"),
        ("RIB (document)", "Fourni" if contrat.rib_document else "NON FOURNI"),
    ]))

    beneficiaires = [
        (f"Bénéficiaire {i}", _beneficiaire_label(b))
        for i, b in enumerate(contrat.beneficiaires.all(), 1)
    ]
    if beneficiaires:
        sections.append(("BÉNÉFICIAIRES", beneficiaires))

    correspondants = [
        (f"Correspondant {c.numero}", f"{c.nom} {c.prenom} — {c.telephone}")
        for c in contrat.correspondants.all()
    ]
    if correspondants:
        sections.append(("CORRESPONDANTS", correspondants))

    sections.append(("CONSEILLER / AGENT", _conseiller_rows(contrat)))
    return sections


def _prevoyance_sections(contrat):
    sections = [
        ("IDENTIFIANTS", [
            ("N° Police", _fmt(contrat.num_police) if contrat.num_police else "En attente"),
            ("N° Client", _fmt(contrat.num_client)),
            ("Code Inspection", _fmt(contrat.code_inspection)),
            ("Statut", contrat.get_statut_display()),
            ("Date de souscription", _fmt(contrat.created_at.date() if contrat.created_at else None)),
            ("Date d'effet", _fmt(contrat.date_effet)),
            ("Date d'échéance", _fmt(contrat.date_echeance)),
        ]),
        ("SOUSCRIPTEUR", [
            ("Nom complet", f"{contrat.client.nom} {contrat.client.prenoms}"),
            ("Civilité", _fmt(contrat.civilite)),
            ("Date de naissance", _fmt(contrat.date_naissance)),
            ("NNI", _fmt(contrat.nni)),
            ("Type pièce", contrat.get_type_piece_display() if contrat.type_piece else "—"),
            ("N° pièce", _fmt(contrat.numero_piece)),
            ("Situation matrimoniale", contrat.get_situation_matrimoniale_display() if contrat.situation_matrimoniale else "—"),
            ("Téléphone", _fmt(contrat.telephone)),
            ("WhatsApp", _fmt(contrat.whatsapp)),
            ("Email personnel", _fmt(contrat.email_personnel)),
            ("Adresse postale", _fmt(contrat.adresse_postale)),
            ("Profession", _fmt(contrat.profession_detail)),
        ]),
        ("GARANTIES", [
            ("Prévoyance Décès (45 000 FCFA/an)", _fmt(contrat.garantie_prevoyance_deces)),
            ("Frais Funéraires (58 000 FCFA/an)", _fmt(contrat.garantie_frais_funeraires)),
            ("Prime annuelle totale", _money(contrat.prime_annuelle)),
            ("Frais fixes", _money(contrat.frais_fixes)),
        ]),
    ]

    assures = [
        (
            a.get_lien_parente_display(),
            f"{a.nom} — né(e) le {_fmt(a.date_naissance)}"
            + (f" — Contact : {a.contact}" if a.contact else "")
            + f" — Indemnité : {_money(a.montant_indemnite)}",
        )
        for a in contrat.assures_associes.all()
    ]
    if assures:
        sections.append(("ASSURÉS ASSOCIÉS", assures))

    beneficiaires = [
        (f"Bénéficiaire {i}", _beneficiaire_label(b))
        for i, b in enumerate(contrat.beneficiaires.all(), 1)
    ]
    if beneficiaires:
        sections.append(("BÉNÉFICIAIRES", beneficiaires))

    if contrat.corresp_nom:
        sections.append(("CORRESPONDANT", [
            ("Nom", f"{contrat.corresp_nom} {contrat.corresp_prenom}"),
            ("Téléphone", _fmt(contrat.corresp_telephone)),
            ("Email", _fmt(contrat.corresp_email)),
        ]))

    sections.append(("CONSEILLER / AGENT", _conseiller_rows(contrat)))
    return sections


# ═══════════════════════════════════════════════════════════════════════════
#  EXCEL
# ═══════════════════════════════════════════════════════════════════════════

_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
_SECTION_FILL = PatternFill("solid", fgColor="2E75B6")
_ALT_FILL = PatternFill("solid", fgColor="EBF3FB")
_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def _xl_title(ws, text, row):
    ws.merge_cells(f"A{row}:B{row}")
    cell = ws[f"A{row}"]
    cell.value = text
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = _SECTION_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cell.border = _BORDER
    ws.row_dimensions[row].height = 22
    return row + 1


def _xl_row(ws, label, value, row, alt=False):
    a = ws[f"A{row}"]
    b = ws[f"B{row}"]
    a.value = label
    b.value = value
    a.font = Font(bold=True, size=10)
    b.font = Font(size=10)
    if alt:
        a.fill = _ALT_FILL
        b.fill = _ALT_FILL
    a.border = _BORDER
    b.border = _BORDER
    a.alignment = Alignment(indent=1)
    b.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[row].height = 18
    return row + 1


def _export_excel(title_text, sheet_title, filename, sections, num_police):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 55

    ws.merge_cells("A1:B1")
    title = ws["A1"]
    title.value = title_text
    title.font = Font(bold=True, color="FFFFFF", size=14)
    title.fill = _HEADER_FILL
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:B2")
    sub = ws["A2"]
    sub.value = f"Exporté le {date.today().strftime('%d/%m/%Y')}   |   N° Police : {num_police or 'En attente'}"
    sub.font = Font(italic=True, size=10, color="555555")
    sub.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 18

    r = 4
    for section_title, rows in sections:
        r = _xl_title(ws, section_title, r)
        alt = False
        for lbl, val in rows:
            alt = not alt
            r = _xl_row(ws, lbl, val, r, alt)
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_epargne_excel(contrat):
    return _export_excel(
        "MAGIST'ASSUR — Contrat Épargne",
        "Contrat Épargne",
        f"epargne_{contrat.id}_{date.today().isoformat()}.xlsx",
        _epargne_sections(contrat),
        contrat.num_police,
    )


def export_prevoyance_excel(contrat):
    return _export_excel(
        "MAGIST'ASSUR — Contrat Prévoyance",
        "Contrat Prévoyance",
        f"prevoyance_{contrat.id}_{date.today().isoformat()}.xlsx",
        _prevoyance_sections(contrat),
        contrat.num_police,
    )


# ═══════════════════════════════════════════════════════════════════════════
#  PDF  (reportlab)
# ═══════════════════════════════════════════════════════════════════════════

def _pdf_section(elements, text, styles):
    from reportlab.platypus import Paragraph, Spacer
    from reportlab.lib.units import mm
    elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph(text, styles["ma_section"]))
    elements.append(Spacer(1, 2 * mm))


def _pdf_rows(elements, rows, styles):
    from reportlab.platypus import Table, TableStyle, Paragraph
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    # Valeurs saisies par le client : échappées, sinon ReportLab interprète le
    # balisage (<img src="/chemin/serveur"> lirait un fichier du disque).
    data = [[lbl, Paragraph(xml_escape(str(val)), styles["ma_cell"])] for lbl, val in rows]
    t = Table(data, colWidths=[70 * mm, 100 * mm])
    style = TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.Color(0.93, 0.96, 0.98)]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.Color(0.8, 0.8, 0.8)),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ])
    t.setStyle(style)
    elements.append(t)


def _build_pdf_styles():
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ma_title",
        fontName="Helvetica-Bold",
        fontSize=16,
        textColor=colors.Color(0.12, 0.23, 0.37),
        alignment=TA_CENTER,
        spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        name="ma_subtitle",
        fontName="Helvetica",
        fontSize=10,
        textColor=colors.grey,
        alignment=TA_CENTER,
        spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        name="ma_section",
        fontName="Helvetica-Bold",
        fontSize=10,
        textColor=colors.white,
        backColor=colors.Color(0.18, 0.46, 0.71),
        leftIndent=4,
        spaceBefore=2,
        spaceAfter=2,
        leading=16,
    ))
    styles.add(ParagraphStyle(
        name="ma_cell",
        fontName="Helvetica",
        fontSize=9,
    ))
    return styles


def _export_pdf(title_text, doc_title, filename, sections, num_police):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib import colors

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm,
        title=doc_title,
    )

    styles = _build_pdf_styles()
    elements = []

    elements.append(Paragraph(title_text, styles["ma_title"]))
    elements.append(Paragraph(
        f"N° Police : <b>{xml_escape(num_police or 'En attente')}</b>   |   "
        f"Exporté le {date.today().strftime('%d/%m/%Y')}",
        styles["ma_subtitle"]
    ))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.Color(0.18, 0.46, 0.71)))
    elements.append(Spacer(1, 4 * mm))

    for section_title, rows in sections:
        _pdf_section(elements, section_title, styles)
        _pdf_rows(elements, rows, styles)

    doc.build(elements)
    buf.seek(0)
    response = HttpResponse(buf.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_epargne_pdf(contrat):
    return _export_pdf(
        "MAGIST'ASSUR — Contrat Épargne",
        f"Contrat Épargne — {contrat.num_police or 'En attente'}",
        f"epargne_{contrat.id}_{date.today().isoformat()}.pdf",
        _epargne_sections(contrat),
        contrat.num_police,
    )


def export_prevoyance_pdf(contrat):
    return _export_pdf(
        "MAGIST'ASSUR — Contrat Prévoyance",
        f"Contrat Prévoyance — {contrat.num_police or 'En attente'}",
        f"prevoyance_{contrat.id}_{date.today().isoformat()}.pdf",
        _prevoyance_sections(contrat),
        contrat.num_police,
    )
