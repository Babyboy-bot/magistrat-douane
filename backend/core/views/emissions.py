import csv
import io
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.db.models import DecimalField, ExpressionWrapper, F, Q, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import generics, permissions, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response
from rest_framework.views import APIView

from core.models import Emission, Paiement
from core.permissions import IsAgentOrAdmin
from core.serializers import EmissionSerializer, PaiementSerializer
from core.services import create_emission, create_paiement


class EmissionPagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = "page_size"
    max_page_size = 200


def _emissions_version_key(user_id: int) -> str:
    return f"emissions_ver_{user_id}"


def _get_emissions_version(user_id: int) -> int:
    v = cache.get(_emissions_version_key(user_id))
    if v is None:
        cache.set(_emissions_version_key(user_id), 1, timeout=None)
        return 1
    return v


def invalidate_emissions_cache(user_id: int) -> None:
    """Incrémente la version → toutes les entrées précédentes deviennent orphelines."""
    key = _emissions_version_key(user_id)
    try:
        cache.incr(key)
    except ValueError:
        cache.set(key, 1, timeout=None)


def invalidate_admin_emissions_cache() -> None:
    """Invalide le cache émissions pour tous les admins actifs."""
    from django.contrib.auth import get_user_model
    User = get_user_model()
    admin_ids = User.objects.filter(role="admin", is_active=True).values_list("id", flat=True)
    for admin_id in admin_ids:
        invalidate_emissions_cache(admin_id)


class EmissionListCreateView(generics.ListCreateAPIView):
    serializer_class = EmissionSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = EmissionPagination

    def list(self, request, *args, **kwargs):
        ver = _get_emissions_version(request.user.id)
        params = request.query_params.urlencode()
        cache_key = f"emissions_{request.user.id}_v{ver}_{params}"
        cached = cache.get(cache_key)
        if cached is not None:
            return Response(cached)
        response = super().list(request, *args, **kwargs)
        cache.set(cache_key, response.data, timeout=300)
        return response

    filterset_fields = {
        "client": ["exact"],
        "agent": ["exact"],
        "compagnie": ["exact", "icontains"],
        "risque": ["exact", "icontains"],
        "situation": ["exact"],
        "date": ["exact", "gte", "lte"],
        "echeance": ["exact", "gte", "lte"],
    }
    search_fields = [
        "client__nom", "client__prenoms",
        "num_police", "num_facture",
        "immatriculation", "marque", "compagnie", "risque",
    ]
    ordering_fields = ["date", "echeance", "prime_ttc", "prime_nette", "created_at"]
    ordering = ["-date"]

    def get_permissions(self):
        if self.request.method == "POST":
            return [permissions.IsAuthenticated(), IsAgentOrAdmin()]
        return [permissions.IsAuthenticated()]

    def get_queryset(self):
        user = self.request.user
        qs = Emission.objects.select_related("client", "agent").prefetch_related("paiements")
        # Admin : portefeuille complet du cabinet
        if user.role == "admin":
            return qs
        # Agent : uniquement ses propres émissions (isolation RGPD)
        if user.role == "agent":
            return qs.filter(agent=user)
        # Client : uniquement ses propres contrats
        return qs.filter(client=user)

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            emission = create_emission(
                validated_data=serializer.validated_data,
                agent=request.user,
            )
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        invalidate_emissions_cache(request.user.id)
        if emission.client_id:
            invalidate_emissions_cache(emission.client_id)
        invalidate_admin_emissions_cache()
        return Response(EmissionSerializer(emission).data, status=status.HTTP_201_CREATED)


class EmissionDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = EmissionSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        qs = Emission.objects.select_related("client", "agent").prefetch_related("paiements")
        if user.role == "admin":
            return qs
        if user.role == "agent":
            return qs.filter(agent=user)
        return qs.filter(client=user)

    def _check_write_permission(self):
        if self.request.user.role not in ("agent", "admin"):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("Seuls les agents et administrateurs peuvent modifier ou supprimer une émission.")

    def update(self, request, *args, **kwargs):
        self._check_write_permission()
        emission = self.get_object()
        client_id = emission.client_id
        response = super().update(request, *args, **kwargs)
        invalidate_emissions_cache(request.user.id)
        if client_id:
            invalidate_emissions_cache(client_id)
        invalidate_admin_emissions_cache()
        return response

    def destroy(self, request, *args, **kwargs):
        self._check_write_permission()
        emission = self.get_object()
        client_id = emission.client_id
        response = super().destroy(request, *args, **kwargs)
        invalidate_emissions_cache(request.user.id)
        if client_id:
            invalidate_emissions_cache(client_id)
        invalidate_admin_emissions_cache()
        return response


class PaiementCreateView(generics.CreateAPIView):
    serializer_class = PaiementSerializer
    permission_classes = [permissions.IsAuthenticated, IsAgentOrAdmin]

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = dict(serializer.validated_data)
        emission = data.pop("emission")
        if request.user.role == "agent" and emission.agent_id != request.user.id:
            return Response(
                {"detail": "Vous ne pouvez enregistrer un paiement que sur vos propres émissions."},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            paiement = create_paiement(
                validated_data=data,
                agent=request.user,
                emission=emission,
            )
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        invalidate_emissions_cache(request.user.id)
        if emission.client_id:
            invalidate_emissions_cache(emission.client_id)
        return Response(PaiementSerializer(paiement).data, status=status.HTTP_201_CREATED)


@api_view(["GET"])
@permission_classes([permissions.IsAuthenticated])
def emission_totals_view(request):
    user = request.user
    if user.role == "admin":
        qs = Emission.objects.all()
    elif user.role == "agent":
        qs = Emission.objects.filter(agent=user)
    else:
        qs = Emission.objects.filter(client=user)
    p = request.query_params
    if p.get("client"): qs = qs.filter(client=p["client"])
    if p.get("agent"): qs = qs.filter(agent=p["agent"])
    if p.get("compagnie"): qs = qs.filter(compagnie=p["compagnie"])
    if p.get("risque"): qs = qs.filter(risque=p["risque"])
    if p.get("situation"): qs = qs.filter(situation=p["situation"])
    if p.get("date__gte"): qs = qs.filter(date__gte=p["date__gte"])
    if p.get("date__lte"): qs = qs.filter(date__lte=p["date__lte"])
    if p.get("echeance__gte"): qs = qs.filter(echeance__gte=p["echeance__gte"])
    if p.get("echeance__lte"): qs = qs.filter(echeance__lte=p["echeance__lte"])
    if p.get("search"):
        s = p["search"]
        qs = qs.filter(
            Q(client__nom__icontains=s) | Q(client__prenoms__icontains=s) |
            Q(num_police__icontains=s) | Q(num_facture__icontains=s) |
            Q(immatriculation__icontains=s) | Q(marque__icontains=s) |
            Q(compagnie__icontains=s) | Q(risque__icontains=s)
        )
    totals = qs.aggregate(
        prime_nette=Coalesce(Sum("prime_nette"), Value(0, output_field=DecimalField())),
        prime_ttc=Coalesce(Sum("prime_ttc"), Value(0, output_field=DecimalField())),
        total_paye=Coalesce(Sum("paiements__montant_regle"), Value(0, output_field=DecimalField())),
        commission=Coalesce(
            Sum(
                ExpressionWrapper(
                    F("prime_nette") * F("taux_commission") / 100,
                    output_field=DecimalField(),
                )
            ),
            Value(0, output_field=DecimalField()),
        ),
    )
    totals["solde"] = totals["prime_ttc"] - totals["total_paye"]
    totals["count"] = qs.count()
    return Response(totals)


@api_view(["GET"])
@permission_classes([permissions.IsAuthenticated])
def emission_balance_view(request, pk):
    emission = get_object_or_404(Emission, pk=pk)
    user = request.user
    if user.role not in ("agent", "admin") and emission.client != user:
        return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)
    if user.role == "agent" and emission.agent != user:
        return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)
    total_paye = emission.paiements.aggregate(total=Sum("montant_regle"))["total"] or 0
    return Response({
        "emission_id": emission.pk,
        "prime_ttc": emission.prime_ttc,
        "total_paye": total_paye,
        "solde": emission.prime_ttc - total_paye,
    })


@api_view(["GET"])
@permission_classes([permissions.IsAuthenticated])
def export_emissions(request):
    if request.user.role not in ("agent", "admin"):
        return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)

    qs = Emission.objects.select_related("client").prefetch_related("paiements")
    if request.user.role == "agent":
        qs = qs.filter(agent=request.user)
    p = request.query_params
    if p.get("compagnie"): qs = qs.filter(compagnie__icontains=p["compagnie"])
    if p.get("risque"): qs = qs.filter(risque__icontains=p["risque"])
    if p.get("situation"): qs = qs.filter(situation__iexact=p["situation"])
    if p.get("echeance_before"):
        try:
            qs = qs.filter(echeance__lte=datetime.strptime(p["echeance_before"], "%Y-%m-%d").date())
        except ValueError:
            return Response({"detail": "Format de date invalide (YYYY-MM-DD)."}, status=400)
    if p.get("echeance_within_days"):
        try:
            days = int(p["echeance_within_days"])
            if days < 0: raise ValueError
            qs = qs.filter(echeance__lte=timezone.now().date() + timedelta(days=days))
        except ValueError:
            return Response({"detail": "echeance_within_days doit être un entier positif."}, status=400)

    if qs.count() > 10000:
        raise ValidationError("Export limité à 10 000 lignes. Affinez vos filtres.")

    qs = qs.annotate(
        total_paye_value=Coalesce(Sum("paiements__montant_regle"), Value(0, output_field=DecimalField(max_digits=12, decimal_places=2)))
    )
    if p.get("impayes_only", "").lower() in {"1", "true", "yes"}:
        qs = qs.filter(prime_ttc__gt=F("total_paye_value"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Emissions"
    headers = ["N° Police", "Souscripteur", "Compagnie", "Risque", "Prime TTC", "Total Payé", "Solde Restant", "Date Effet", "Situation", "Date Création"]
    ws.append(headers)
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="4B5563")
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.font = hf; c.fill = hfill; c.alignment = Alignment(horizontal="center", vertical="center")

    for e in qs:
        tp = getattr(e, "total_paye_value", None) or e.total_paye or 0
        cat = e.created_at
        if timezone.is_aware(cat): cat = timezone.localtime(cat).replace(tzinfo=None)
        ws.append([e.num_police, e.client.full_name if e.client else "", e.compagnie, e.risque, float(e.prime_ttc), float(tp), float(e.prime_ttc - tp), e.effet, e.situation, cat])

    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for c in row: c.number_format = r"#,##0\ FCFA"
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for c in row: c.number_format = "DD/MM/YYYY"
    for row in ws.iter_rows(min_row=2, min_col=10, max_col=10):
        for c in row: c.number_format = "DD/MM/YYYY HH:MM"
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 2

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    ts = timezone.now().strftime("%Y%m%d_%H%M%S")
    resp = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="emissions_export_{ts}.xlsx"'
    return resp


IMPORT_COLUMNS = ["num_police", "num_facture", "client_email", "date", "risque", "compagnie", "marque", "immatriculation", "effet", "echeance", "prime_nette", "prime_ttc", "situation"]


class EmissionImportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        if request.user.role not in ("agent", "admin"):
            return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)
        file = request.FILES.get("file")
        if not file:
            return Response({"detail": "Aucun fichier fourni"}, status=status.HTTP_400_BAD_REQUEST)
        ext = file.name.rsplit(".", 1)[-1].lower()
        if ext == "csv":
            rows = self._parse_csv(file)
        elif ext in ("xlsx", "xls"):
            rows = self._parse_xlsx(file)
        else:
            return Response({"detail": "Format non supporté. Utilisez CSV ou XLSX."}, status=400)

        User = get_user_model()
        created, skipped, errors = 0, 0, []
        for i, row in enumerate(rows, start=2):
            line_errors = self._validate_row(row, i)
            if line_errors: errors.extend(line_errors); continue
            try:
                client = User.objects.get(email__iexact=row["client_email"].strip())
            except User.DoesNotExist:
                errors.append(f"Ligne {i} : client introuvable ({row['client_email']})"); continue
            if Emission.objects.filter(num_police=row["num_police"].strip()).exists():
                skipped += 1; continue
            try:
                Emission.objects.create(
                    client=client, agent=request.user,
                    date=self._parse_date(row["date"]), risque=row["risque"].strip(),
                    compagnie=row["compagnie"].strip(), num_facture=row["num_facture"].strip(),
                    num_police=row["num_police"].strip(),
                    marque=row.get("marque", "-").strip() or "-",
                    immatriculation=row.get("immatriculation", "-").strip() or "-",
                    effet=self._parse_date(row["effet"]), echeance=self._parse_date(row["echeance"]),
                    prime_nette=int(Decimal(str(row["prime_nette"]).replace(" ", "").replace(",", "."))),
                    prime_ttc=int(Decimal(str(row["prime_ttc"]).replace(" ", "").replace(",", "."))),
                    situation=row.get("situation", "Affaire nouvelle").strip() or "Affaire nouvelle",
                )
                created += 1
            except Exception as exc:
                errors.append(f"Ligne {i} : {exc}")

        if created > 0:
            invalidate_emissions_cache(request.user.id)
            invalidate_admin_emissions_cache()
        return Response({"created": created, "skipped": skipped, "errors": errors, "total_rows": created + skipped + len(errors)})

    def _parse_csv(self, file):
        content = file.read().decode("utf-8-sig")
        import csv as _csv
        sample = content[:2048]
        try:
            dialect = _csv.Sniffer().sniff(sample, delimiters=";,")
        except _csv.Error:
            dialect = None
        if dialect:
            return list(_csv.DictReader(io.StringIO(content), dialect=dialect))
        return list(_csv.DictReader(io.StringIO(content), delimiter=";"))

    def _parse_xlsx(self, file):
        from openpyxl import load_workbook
        wb = load_workbook(filename=file, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        if not rows: return []
        headers = [str(h).strip() if h else "" for h in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(c is not None for c in row)]

    def _validate_row(self, row, line_num):
        errors = []
        for f in ["num_police", "num_facture", "client_email", "date", "risque", "compagnie", "effet", "echeance", "prime_nette", "prime_ttc"]:
            if not row.get(f): errors.append(f"Ligne {line_num} : champ '{f}' manquant")
        if errors: return errors
        try:
            pn = float(str(row["prime_nette"]).replace(" ", "").replace(",", "."))
            pt = float(str(row["prime_ttc"]).replace(" ", "").replace(",", "."))
            if pn > pt: errors.append(f"Ligne {line_num} : prime_nette supérieure à prime_ttc")
        except (ValueError, KeyError):
            errors.append(f"Ligne {line_num} : prime_nette ou prime_ttc invalides")
        try:
            if self._parse_date(str(row["echeance"])) <= self._parse_date(str(row["effet"])):
                errors.append(f"Ligne {line_num} : échéance doit être postérieure à l'effet")
        except ValueError:
            pass
        return errors

    def _parse_date(self, value):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try: return datetime.strptime(str(value).strip(), fmt).date()
            except ValueError: continue
        raise ValueError(f"Date invalide : {value}")
