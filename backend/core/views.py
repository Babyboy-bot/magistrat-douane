import csv
import io
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from datetime import datetime, timedelta

from django.db.models import DecimalField, F, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import generics, permissions, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response
from rest_framework.views import APIView

from notifications.emails import (
    send_request_acknowledgment,
    send_new_request_alert_to_agents,
    send_request_treated,
    send_livrable_notification,
)
from notifications.ws import (
    notify_agents_new_request,
    notify_client_request_treated,
    notify_client_livrable_added,
)
from .models import Emission, Paiement, ClientRequest, Document
from .serializers import (
    EmissionSerializer,
    PaiementSerializer,
    ClientRequestSerializer,
    ClientRequestCreateSerializer,
    DocumentSerializer,
)


class IsAgentOrAdmin(permissions.BasePermission):
    def has_permission(self, request, view):
        return request.user.role in ("agent", "admin")


# ── Emissions ──

class EmissionPagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = "page_size"
    max_page_size = 200


class EmissionListCreateView(generics.ListCreateAPIView):
    serializer_class = EmissionSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = EmissionPagination
    filterset_fields = {
        "client": ["exact"],
        "agent": ["exact"],
        "compagnie": ["exact", "icontains"],
        "risque": ["exact", "icontains"],
        "situation": ["exact"],
        "date": ["exact", "gte", "lte"],
        "echeance": ["exact", "gte", "lte"],
    }
    search_fields = [
        "client__nom",
        "client__prenoms",
        "num_police",
        "num_facture",
        "immatriculation",
        "marque",
        "compagnie",
        "risque",
    ]
    ordering_fields = ["date", "echeance", "prime_ttc", "prime_nette", "created_at"]
    ordering = ["-date"]

    def get_queryset(self):
        user = self.request.user
        if user.role in ("agent", "admin"):
            return Emission.objects.all()
        return Emission.objects.filter(client=user)

    def perform_create(self, serializer):
        serializer.save(agent=self.request.user)


class EmissionDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = EmissionSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.role in ("agent", "admin"):
            return Emission.objects.all()
        return Emission.objects.filter(client=user)


# ── Paiements ──

class PaiementCreateView(generics.CreateAPIView):
    serializer_class = PaiementSerializer
    permission_classes = [permissions.IsAuthenticated, IsAgentOrAdmin]

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)


# ── Emission balance endpoint ──

@api_view(["GET"])
@permission_classes([permissions.IsAuthenticated])
def emission_balance_view(request, pk):
    emission = get_object_or_404(Emission, pk=pk)
    user = request.user
    if user.role not in ("agent", "admin") and emission.client != user:
        return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)

    total_paye = emission.paiements.aggregate(total=Sum("montant_regle"))["total"] or 0
    solde = emission.prime_ttc - total_paye

    return Response(
        {
            "emission_id": emission.pk,
            "prime_ttc": emission.prime_ttc,
            "total_paye": total_paye,
            "solde": solde,
        }
    )


# ── Requests ──

class ClientRequestListCreateView(generics.ListCreateAPIView):
    permission_classes = [permissions.IsAuthenticated]

    def get_serializer_class(self):
        if self.request.method == "POST":
            return ClientRequestCreateSerializer
        return ClientRequestSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role in ("agent", "admin"):
            return ClientRequest.objects.all()
        if user.account_status in ("active", "verified"):
            return ClientRequest.objects.filter(client=user)
        return ClientRequest.objects.none()

    def perform_create(self, serializer):
        instance = serializer.save()
        send_request_acknowledgment(instance)
        send_new_request_alert_to_agents(instance)
        notify_agents_new_request(instance)


class ClientRequestDetailView(generics.RetrieveUpdateAPIView):
    serializer_class = ClientRequestSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.role in ("agent", "admin"):
            return ClientRequest.objects.all()
        # Les utilisateurs actifs peuvent voir leurs propres requêtes
        if user.account_status in ("active", "verified"):
            return ClientRequest.objects.filter(client=user)
        return ClientRequest.objects.none()


@api_view(["POST"])
@permission_classes([permissions.IsAuthenticated])
def treat_request_view(request, pk):
    if request.user.role not in ("agent", "admin"):
        return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)
    try:
        req = ClientRequest.objects.get(pk=pk)
    except ClientRequest.DoesNotExist:
        return Response({"detail": "Requête introuvable"}, status=status.HTTP_404_NOT_FOUND)

    # Un agent ne peut traiter que sa propre requête (sauf admin)
    if request.user.role == "agent" and req.agent_id and req.agent_id != request.user.id:
        return Response(
            {"detail": "Cette requête est prise en charge par un autre agent."},
            status=status.HTTP_403_FORBIDDEN,
        )

    req.status = "traite"
    req.treated_at = timezone.now()
    req.agent = request.user
    req.save()
    send_request_treated(req)
    notify_client_request_treated(req)
    return Response(ClientRequestSerializer(req).data)


@api_view(["POST"])
@permission_classes([permissions.IsAuthenticated])
def claim_request_view(request, pk):
    """Un agent prend en charge une requête disponible."""
    if request.user.role not in ("agent", "admin"):
        return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)

    from django.db import transaction
    try:
        with transaction.atomic():
            req = ClientRequest.objects.select_for_update().get(pk=pk)
            if req.status != ClientRequest.Status.EN_ATTENTE:
                return Response(
                    {"detail": "Cette requête est déjà prise en charge ou traitée."},
                    status=status.HTTP_409_CONFLICT,
                )
            req.agent = request.user
            req.status = ClientRequest.Status.EN_COURS
            req.save()
    except ClientRequest.DoesNotExist:
        return Response({"detail": "Requête introuvable"}, status=status.HTTP_404_NOT_FOUND)

    # Notifier le client
    from notifications.service import notification_service
    notification_service.send_notification_to_user(
        user_id=req.client_id,
        title="Votre requête est en cours de traitement",
        message=(
            f"Votre demande {req.numero} ({req.get_type_display()}) "
            f"a été prise en charge par {request.user.full_name}."
        ),
        type="info",
        priority="medium",
        data={"request_id": req.id, "numero": req.numero},
    )
    return Response(ClientRequestSerializer(req).data)


@api_view(["POST"])
@permission_classes([permissions.IsAuthenticated])
def release_request_view(request, pk):
    """Un agent libère une requête qu'il avait prise en charge."""
    if request.user.role not in ("agent", "admin"):
        return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)
    try:
        req = ClientRequest.objects.get(pk=pk)
    except ClientRequest.DoesNotExist:
        return Response({"detail": "Requête introuvable"}, status=status.HTTP_404_NOT_FOUND)

    if req.status != ClientRequest.Status.EN_COURS:
        return Response(
            {"detail": "Cette requête n'est pas en cours de traitement."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    if request.user.role == "agent" and req.agent_id != request.user.id:
        return Response(
            {"detail": "Vous ne pouvez libérer que vos propres requêtes."},
            status=status.HTTP_403_FORBIDDEN,
        )

    req.agent = None
    req.status = ClientRequest.Status.EN_ATTENTE
    req.save()
    return Response(ClientRequestSerializer(req).data)


# ── Documents ──

class RequestLivrablesView(APIView):
    """
    GET  : liste les livrables d'une requête (client propriétaire ou agent/admin)
    POST : l'agent upload un livrable sur la requête
    """
    permission_classes = [permissions.IsAuthenticated]

    def _get_request_or_403(self, pk, user):
        req = get_object_or_404(ClientRequest, pk=pk)
        if user.role not in ("agent", "admin") and req.client != user:
            return None, Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)
        return req, None

    def get(self, request, pk):
        req, err = self._get_request_or_403(pk, request.user)
        if err:
            return err
        livrables = req.documents.filter(is_livrable=True)
        return Response(DocumentSerializer(livrables, many=True, context={"request": request}).data)

    def post(self, request, pk):
        if request.user.role not in ("agent", "admin"):
            return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)
        req, err = self._get_request_or_403(pk, request.user)
        if err:
            return err
        serializer = DocumentSerializer(data=request.data, context={"request": request})
        serializer.is_valid(raise_exception=True)
        doc = serializer.save(
            request=req,
            uploaded_by=request.user,
            is_livrable=True,
        )
        notify_client_livrable_added(req, doc)
        send_livrable_notification(req, doc)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class DocumentUploadView(generics.CreateAPIView):
    serializer_class = DocumentSerializer

    def perform_create(self, serializer):
        serializer.save(uploaded_by=self.request.user)


class CarteGriseUploadView(generics.CreateAPIView):
    """Upload simple des cartes grises sans OCR"""
    serializer_class = DocumentSerializer
    permission_classes = [permissions.IsAuthenticated]

    def create(self, request, *args, **kwargs):
        try:
            file = request.FILES.get('file')
            if not file:
                return Response(
                    {"error": "Aucun fichier fourni"}, 
                    status=status.HTTP_400_BAD_REQUEST
                )

            allowed_types = ['application/pdf', 'image/jpeg', 'image/jpg', 'image/png']
            if file.content_type not in allowed_types:
                return Response(
                    {"error": "Format de fichier non supporté (PDF, JPG, PNG uniquement)"}, 
                    status=status.HTTP_400_BAD_REQUEST
                )

            # Création du document
            document = Document.objects.create(
                nom=f"carte_grise_{file.name}",
                fichier=file,
                type_fichier="carte_grise",
                uploaded_by=request.user,
                is_livrable=False
            )
            
            return Response({
                "document_id": document.id,
                "document_nom": document.nom,
                "file_size": file.size,
                "file_type": file.content_type,
                "message": "Carte grise uploadée avec succès"
            })

        except Exception as e:
            return Response(
                {"error": f"Erreur lors du traitement: {str(e)}"}, 
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


# ── Compagnies ──

COMPAGNIES_DATA = [
    "Activa Assurances",
    "AFG Assurances",
    "Amsa Assurances",
    "Atlanta Assurances",
    "Atlantique Assurances",
    "AXA Assurances",
    "Coris Assurances",
    "GNA CI",
    "NSIA Assurances",
    "SanlamAllianz",
    "SCHIBA",
    "Sérenity",
    "SAAR CI",
    "SUNU Assurances IARD",
]


class CompagnieListView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        search = request.query_params.get("search", "").strip().lower()
        results = []
        for name in COMPAGNIES_DATA:
            if search and search not in name.lower():
                continue
            results.append({"id": slugify(name), "name": name})
        return Response(results)


# ── Dashboard Stats ──

@api_view(["GET"])
@permission_classes([permissions.IsAuthenticated])
def dashboard_stats_view(request):
    user = request.user
    if user.role in ("agent", "admin"):
        emissions = Emission.objects.all()
        # Stats utilisateurs pour admin/agent
        from django.contrib.auth import get_user_model
        User = get_user_model()
        total_clients = User.objects.filter(role="client").count()
        total_prospects = User.objects.filter(role="prospect").count()
        total_agents = User.objects.filter(role="agent").count()
    else:
        emissions = Emission.objects.filter(client=user)
        total_clients = 0
        total_prospects = 0
        total_agents = 0

    now = timezone.now().date()
    actifs = emissions.filter(echeance__gt=now)
    expires = emissions.filter(echeance__lte=now)
    total_primes = emissions.aggregate(t=Sum("prime_ttc"))["t"] or 0
    primes_payees = Paiement.objects.filter(
        emission__in=emissions
    ).aggregate(t=Sum("montant_regle"))["t"] or 0

    return Response({
        "total_contrats": emissions.count(),
        "contrats_actifs": actifs.count(),
        "contrats_expires": expires.count(),
        "total_primes": total_primes,
        "primes_payees": primes_payees,
        "primes_impayees": total_primes - primes_payees,
        "total_clients": total_clients,
        "total_prospects": total_prospects,
        "total_agents": total_agents,
    })


@api_view(["GET"])
@permission_classes([permissions.IsAuthenticated])
def export_emissions(request):
    if request.user.role not in ("agent", "admin"):
        return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)

    queryset = Emission.objects.select_related("client").prefetch_related("paiements")

    compagnie = request.query_params.get("compagnie")
    risque = request.query_params.get("risque")
    situation = request.query_params.get("situation")
    echeance_before = request.query_params.get("echeance_before")
    echeance_within_days = request.query_params.get("echeance_within_days")
    impayes_only = request.query_params.get("impayes_only")

    if compagnie:
        queryset = queryset.filter(compagnie__icontains=compagnie)
    if risque:
        queryset = queryset.filter(risque__icontains=risque)
    if situation:
        queryset = queryset.filter(situation__iexact=situation)
    if echeance_before:
        try:
            date_limit = datetime.strptime(echeance_before, "%Y-%m-%d").date()
            queryset = queryset.filter(echeance__lte=date_limit)
        except ValueError:
            return Response({"detail": "Format de date invalide pour echeance_before (YYYY-MM-DD attendu)."}, status=status.HTTP_400_BAD_REQUEST)
    if echeance_within_days:
        try:
            days = int(echeance_within_days)
            if days < 0:
                raise ValueError
            date_limit = timezone.now().date() + timedelta(days=days)
            queryset = queryset.filter(echeance__lte=date_limit)
        except ValueError:
            return Response({"detail": "echeance_within_days doit être un entier positif."}, status=status.HTTP_400_BAD_REQUEST)

    total_count = queryset.count()
    if total_count > 10000:
        raise ValidationError("Export limité à 10 000 lignes. Affinez vos filtres.")

    queryset = queryset.annotate(
        total_paye_value=Coalesce(
            Sum("paiements__montant_regle"),
            Value(0, output_field=DecimalField(max_digits=12, decimal_places=2)),
        )
    )

    if impayes_only and str(impayes_only).lower() in {"1", "true", "yes"}:
        queryset = queryset.filter(prime_ttc__gt=F("total_paye_value"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Emissions"

    headers = [
        "N° Police",
        "Souscripteur",
        "Compagnie",
        "Risque",
        "Prime TTC",
        "Total Payé",
        "Solde Restant",
        "Date Effet",
        "Situation",
        "Date Création",
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4B5563")

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    currency_format = "#,##0 \FCFA"
    date_format = "DD/MM/YYYY"
    datetime_format = "DD/MM/YYYY HH:MM"

    for emission in queryset:
        total_paye = getattr(emission, "total_paye_value", None)
        if total_paye is None:
            total_paye = emission.total_paye or 0
        solde = emission.prime_ttc - total_paye
        created_at = emission.created_at
        if timezone.is_aware(created_at):
            created_at = timezone.localtime(created_at).replace(tzinfo=None)
        ws.append([
            emission.num_police,
            emission.client.full_name if emission.client else "",
            emission.compagnie,
            emission.risque,
            float(emission.prime_ttc),
            float(total_paye),
            float(solde),
            emission.effet,
            emission.situation,
            created_at,
        ])

    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for cell in row:
            cell.number_format = currency_format

    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for cell in row:
            cell.number_format = date_format

    for row in ws.iter_rows(min_row=2, min_col=10, max_col=10):
        for cell in row:
            cell.number_format = datetime_format

    for col in ws.columns:
        max_length = 0
        column = col[0].column
        for cell in col:
            try:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
            except ValueError:
                continue
        adjusted_width = max_length + 2
        ws.column_dimensions[get_column_letter(column)].width = adjusted_width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="emissions_export_{timestamp}.xlsx"'
    return response


# ── Import portefeuille ──

IMPORT_COLUMNS = [
    "num_police", "num_facture", "client_email",
    "date", "risque", "compagnie",
    "marque", "immatriculation",
    "effet", "echeance",
    "prime_nette", "prime_ttc",
    "situation",
]


class EmissionImportView(APIView):
    """
    POST /api/emissions/import/
    Accepte un fichier CSV ou XLSX et crée les émissions ligne par ligne.
    Retourne un rapport {created, skipped, errors}.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        if request.user.role not in ("agent", "admin"):
            return Response({"detail": "Permission refusée"}, status=status.HTTP_403_FORBIDDEN)

        file = request.FILES.get("file")
        if not file:
            return Response({"detail": "Aucun fichier fourni"}, status=status.HTTP_400_BAD_REQUEST)

        ext = file.name.rsplit(".", 1)[-1].lower()
        if ext == "csv":
            rows = self._parse_csv(file)
        elif ext in ("xlsx", "xls"):
            rows = self._parse_xlsx(file)
        else:
            return Response(
                {"detail": "Format non supporté. Utilisez CSV ou XLSX."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        User = get_user_model()
        created, skipped, errors = 0, 0, []

        for i, row in enumerate(rows, start=2):
            line_errors = self._validate_row(row, i)
            if line_errors:
                errors.extend(line_errors)
                continue

            try:
                client = User.objects.get(email__iexact=row["client_email"].strip())
            except User.DoesNotExist:
                errors.append(f"Ligne {i} : client introuvable ({row['client_email']})")
                continue

            if Emission.objects.filter(num_police=row["num_police"].strip()).exists():
                skipped += 1
                continue

            try:
                Emission.objects.create(
                    client=client,
                    agent=request.user,
                    date=self._parse_date(row["date"]),
                    risque=row["risque"].strip(),
                    compagnie=row["compagnie"].strip(),
                    num_facture=row["num_facture"].strip(),
                    num_police=row["num_police"].strip(),
                    marque=row.get("marque", "-").strip() or "-",
                    immatriculation=row.get("immatriculation", "-").strip() or "-",
                    effet=self._parse_date(row["effet"]),
                    echeance=self._parse_date(row["echeance"]),
                    prime_nette=int(float(str(row["prime_nette"]).replace(" ", "").replace(",", "."))),
                    prime_ttc=int(float(str(row["prime_ttc"]).replace(" ", "").replace(",", "."))),
                    situation=row.get("situation", "Affaire nouvelle").strip() or "Affaire nouvelle",
                )
                created += 1
            except Exception as exc:
                errors.append(f"Ligne {i} : {exc}")

        return Response({
            "created": created,
            "skipped": skipped,
            "errors": errors,
            "total_rows": created + skipped + len(errors),
        })

    def _parse_csv(self, file):
        text = file.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text), delimiter=";")
        return list(reader)

    def _parse_xlsx(self, file):
        from openpyxl import load_workbook
        wb = load_workbook(filename=file, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h else "" for h in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(c is not None for c in row)]

    def _validate_row(self, row, line_num):
        errors = []
        required = ["num_police", "num_facture", "client_email", "date",
                    "risque", "compagnie", "effet", "echeance", "prime_nette", "prime_ttc"]
        for field in required:
            if not row.get(field):
                errors.append(f"Ligne {line_num} : champ '{field}' manquant")
        return errors

    def _parse_date(self, value):
        value = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
        raise ValueError(f"Date invalide : {value}")
