"""Audit lot 3 — injection de formules, HTML dans les emails, sonde de santé."""
import csv
import io

import pytest
from django.core import mail
from openpyxl import load_workbook

from core.models import ClientRequest
from core.spreadsheet import csv_safe
from magistassur.models import ContratEpargne
from notifications.emails import (
    send_new_request_alert_to_agents, send_request_acknowledgment, send_welcome_email,
)

PIEGE = '=HYPERLINK("http://evil.example","Cliquez")'


@pytest.fixture
def client_piege(client_user):
    client_user.nom = PIEGE
    client_user.prenoms = '<a href="http://evil.example">Vérifiez votre compte</a>'
    client_user.save()
    return client_user


class TestCsvSafe:
    @pytest.mark.parametrize("val", ["=1+1", "+33", "-2", "@SUM(A1)", "\tx"])
    def test_prefixe(self, val):
        assert csv_safe(val) == "'" + val

    @pytest.mark.parametrize("val", ["Kouassi", "", 42, None, "a=b"])
    def test_inchange(self, val):
        assert csv_safe(val) == val


@pytest.mark.django_db
class TestFormulesNeutralisees:
    def test_export_magistassur_excel(self, api_client, agent, client_piege):
        c = ContratEpargne.objects.create(
            client=client_piege, duree_annees=10, periodicite="mensuelle",
            montant_cotisation=25000, pris_en_charge_par=agent, adresse_postale=PIEGE,
        )
        api_client.force_authenticate(user=agent)
        resp = api_client.get(f"/api/magistassur/epargne/{c.id}/export-excel/")
        assert resp.status_code == 200
        ws = load_workbook(io.BytesIO(resp.content)).active
        cellules = [cell for row in ws.iter_rows() for cell in row if cell.value == PIEGE]
        assert cellules, "la valeur piégée doit apparaître telle quelle"
        assert all(cell.data_type == "s" for cell in cellules)

    def test_export_emissions_excel(self, api_client, agent, emission, client_piege):
        api_client.force_authenticate(user=agent)
        resp = api_client.get("/api/emissions/export/")
        assert resp.status_code == 200
        ws = load_workbook(io.BytesIO(resp.content)).active
        assert not [c for row in ws.iter_rows() for c in row if c.data_type == "f"]

    def test_export_audit_csv(self, api_client, admin_user, client_piege):
        from accounts.audit import log_user_action
        log_user_action(admin_user, client_piege, "activation")
        api_client.force_authenticate(user=admin_user)
        resp = api_client.get("/api/auth/audit-logs/export/")
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))
        cibles = [r[3] for r in rows[1:]]
        assert all(not c.startswith(("=", "+", "-", "@")) for c in cibles)
        assert any(c.startswith("'=HYPERLINK") for c in cibles)


@pytest.mark.django_db
class TestEmailsEchappes:
    @pytest.fixture(autouse=True)
    def _locmem(self, settings):
        settings.EMAIL_BACKEND = "django.core.mail.backends.locmem.EmailBackend"

    def _html(self):
        return mail.outbox[-1].alternatives[0][0]

    def test_bienvenue(self, client_piege):
        send_welcome_email(client_piege)
        html = self._html()
        assert '<a href="http://evil.example">' not in html
        assert "&lt;a href=" in html

    def test_accuse_et_alerte_agents(self, client_piege, agent):
        req = ClientRequest.objects.create(client=client_piege, type="souscription_magist")
        send_request_acknowledgment(req)
        assert '<a href="http://evil.example">' not in self._html()
        send_new_request_alert_to_agents(req)
        assert '<a href="http://evil.example">' not in self._html()


INSCRIPTION = {
    "nom": "Nouveau", "prenoms": "Magistrat", "profession": "Juge",
    "profession_type": "magistrat", "password": "UnMotDePasse!2026",
}


class _ThreadImmediat:
    """Exécute la cible tout de suite (les tests n'attendent pas un vrai thread)."""
    def __init__(self, target, args=(), daemon=None):
        self.target, self.args = target, args

    def start(self):
        self.target(*self.args)


@pytest.mark.django_db
class TestInscriptionSansEnumeration:
    @pytest.fixture(autouse=True)
    def _setup(self, settings, monkeypatch):
        settings.EMAIL_BACKEND = "django.core.mail.backends.locmem.EmailBackend"
        monkeypatch.setattr("accounts.serializers.threading.Thread", _ThreadImmediat)

    def _register(self, api_client, **data):
        return api_client.post("/api/auth/register/", {**INSCRIPTION, **data}, format="json")

    def test_email_et_telephone_pris_meme_message(self, api_client, client_user):
        client_user.telephone = "0700000001"
        client_user.save()
        par_email = self._register(api_client, email=client_user.email, telephone="0700000099")
        par_tel = self._register(api_client, email="libre@test.ci", telephone="0700000001")
        assert par_email.status_code == par_tel.status_code == 400
        assert par_email.data == par_tel.data
        assert "email" not in par_email.data and "telephone" not in par_email.data

    def test_titulaire_prevenu(self, api_client, client_user):
        mail.outbox.clear()
        self._register(api_client, email=client_user.email.upper(), telephone="0700000099")
        assert [m.to for m in mail.outbox] == [[client_user.email]]
        assert "Tentative d'inscription" in mail.outbox[0].subject

    def test_inscription_normale_inchangee(self, api_client):
        resp = self._register(api_client, email="nouveau@test.ci", telephone="0700000099")
        assert resp.status_code == 201

    def test_format_telephone_toujours_valide(self, api_client):
        resp = self._register(api_client, email="nouveau@test.ci", telephone="123")
        assert resp.status_code == 400
        assert "telephone" in resp.data


@pytest.mark.django_db
class TestHealth:
    def test_ok_sans_authentification(self, api_client):
        resp = api_client.get("/api/health/")
        assert resp.status_code == 200
        assert resp.json() == {"status": "ok"}
